import re, sys
from openpyxl import load_workbook

MD = '/root/.claude/uploads/638c839c-c994-5bc3-9b3f-d02587b3f6fc/cc2e851e-SPECTAKEOFF_ALL_IN_ONE.md'
md = open(MD, encoding='utf-8').read()
sec7 = md[md.index('## 7. Hand-verified'):md.index('## 8. Full code')]
ref = []
for line in sec7.splitlines():
    if not line.startswith('| '):
        continue
    c = [x.strip() for x in line.strip().strip('|').split('|')]
    if len(c) < 8 or c[0] in ('#', '---') or not c[0].isdigit():
        continue
    ref.append({'n': int(c[0]), 'sec': c[1], 'type': c[2], 'item': c[3],
                'qty': c[4], 'unit': c[5], 'page': int(c[6]), 'quote': c[7]})
print('reference rows parsed:', len(ref))

def key(text):
    return re.sub(r'[^a-z0-9]+', '', (text or '').lower())

wb = load_workbook(sys.argv[1])
ws = wb['Closeout Schedule']
cols = [c.value for c in ws[2]]
out = []
for r in ws.iter_rows(min_row=3, values_only=True):
    if not any(v not in (None, '') for v in r):
        continue
    out.append(dict(zip(cols, r)))
print('output rows:', len(out))

# match each reference row to an output row by (section, page, quote-key) + qty/unit
unmatched = []
used = set()
for rr in ref:
    hit = None
    for i, o in enumerate(out):
        if i in used:
            continue
        if o['Spec Section'] != rr['sec'] or o['Page'] != rr['page']:
            continue
        if key(o['Spec Language (verbatim)']) != key(rr['quote']):
            continue
        oq = '' if o['Qty'] in (None, '') else str(o['Qty'])
        ou = o['Unit'] or ''
        if oq != rr['qty'] or ou != rr['unit']:
            continue
        if o['Type'] != rr['type']:
            continue
        hit = i
        break
    if hit is None:
        unmatched.append(rr)
    else:
        used.add(hit)

print()
print('=== MATCHED %d / %d reference rows ===' % (len(ref) - len(unmatched), len(ref)))
for rr in unmatched:
    print('MISSING ref #%d  %s  p.%s  %s  %s %s  | %s' % (rr['n'], rr['sec'], rr['page'], rr['type'], rr['qty'], rr['unit'], rr['item'][:60]))
    # show near misses
    for o in out:
        if o['Spec Section'] == rr['sec'] and o['Page'] == rr['page']:
            print('    near: %s %s %s | %s | %.70s' % (o['Type'], o['Qty'], o['Unit'], (o['Item Required'] or '')[:50], o['Spec Language (verbatim)']))

print()
print('=== EXTRA output rows (not in the 67) ===')
for i, o in enumerate(out):
    if i not in used:
        print('EXTRA  %s  p.%s  %s  %s %s | %s' % (o['Spec Section'], o['Page'], o['Type'], o['Qty'], o['Unit'], (o['Item Required'] or '')[:70]))
        print('       quote: %.160s' % o['Spec Language (verbatim)'])

# Checks sheet: R10 conflict p.398 vs p.400
checks = ['|'.join(str(x) for x in r) for r in wb['Checks'].iter_rows(min_row=2, values_only=True)]
r10 = [c for c in checks if 'R10 CONFLICT' in c]
print()
print('=== R10 CONFLICT findings (%d) ===' % len(r10))
for c in r10:
    print(' ', c)
print('R10 p.398/p.400 present:', any('398' in c and '400' in c for c in r10))

cov = [r for r in wb['Div 23-26 Coverage'].iter_rows(min_row=2, values_only=True)]
flag = [r for r in cov if r[0] == '26 28 18']
print()
print('=== 26 28 18 coverage row ===')
for r in flag:
    print(' ', r)
print('coverage sections:', len(cov), '| CHECK flagged:', sum(1 for r in cov if str(r[5]).startswith('CHECK')))

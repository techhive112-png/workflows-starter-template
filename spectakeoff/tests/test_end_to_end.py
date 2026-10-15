import csv
import json
from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

from spectakeoff.cli import main
from spectakeoff.evaluate import evaluate
from spectakeoff.excel import _validate, build_workbook
from spectakeoff.models import REQUIREMENT, Finding, PageRecord, SourceResult
from spectakeoff.ocr import installed_languages
from synthetic import make_drawing

SPEC = """SECTION 23 05 00 - COMMON WORK RESULTS FOR HVAC

A. Provide one spare filter set for each air handling unit.
B. The warranty period for fans shall be two years from substantial completion.
C. Submit three copies of operation and maintenance manuals.
D. Furnish extra materials equal to 2% of installed quantity.
"""
SPEC2 = """SECTION 23 05 00 - COMMON WORK RESULTS FOR HVAC (continued)

E. The warranty period for fans shall be one year from substantial completion.
F. Provide training for the Owner's staff.
"""


def build_pdf(path: Path, tmp: Path) -> None:
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((50, 60), SPEC, fontsize=10)
    p = doc.new_page()
    p.insert_text((50, 60), SPEC2, fontsize=10)
    # vector drawing sheet: many line paths + short labels
    p = doc.new_page(width=1224, height=792)
    for i in range(320):
        p.draw_line((40 + (i % 40) * 28, 60 + (i // 40) * 80), (60 + (i % 40) * 28, 110 + (i // 40) * 80))
    for x, label in [(80, "AHU-1"), (300, "AHU-2"), (520, "W12x26 BEAM"), (740, "(4) 3/4\" A325 BOLTS")]:
        p.insert_text((x, 720), label, fontsize=9)
    # schedule with ruling lines
    p = doc.new_page()
    rows = [("TAG", "DESCRIPTION", "QTY"), ("AHU-1", "Air handling unit", "1"), ("AHU-2", "Air handling unit", "1"),
            ("EF-1", "Exhaust fan", "2")]
    x0, y0, cw, rh = 60, 80, 150, 24
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            p.insert_text((x0 + c * cw + 4, y0 + r * rh + 16), val, fontsize=9)
    for r in range(len(rows) + 1):
        p.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
    for c in range(4):
        p.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + len(rows) * rh))
    # scanned (raster-only) drawing page
    img_path = tmp / "drawing.png"
    make_drawing(img_path)
    p = doc.new_page(width=1000, height=550)
    p.insert_image(p.rect, filename=str(img_path))
    doc.save(path)


@pytest.mark.skipif("eng" not in installed_languages(), reason="Tesseract eng not installed")
def test_pdf_end_to_end(tmp_path):
    pdf = tmp_path / "project.pdf"
    build_pdf(pdf, tmp_path)
    counts = tmp_path / "counts.csv"
    counts.write_text("item,count,verified_by,source\nair handling unit,2,J. Estimator,M-101 schedule\n")
    out = tmp_path / "out" / "project.xlsx"
    assert main([str(pdf), "-o", str(out), "--counts", str(counts), "--dpi", "200"]) == 0

    data = json.loads(out.with_suffix(".findings.json").read_text())["findings"]
    by_kind = lambda k: [f for f in data if f["kind"] == k]
    reqs = by_kind("requirement")
    texts = [f["text"] for f in reqs]
    assert "A. Provide one spare filter set for each air handling unit." in texts
    pct = next(f for f in reqs if "2%" in f["text"])
    assert (pct["value"], pct["unit"]) == ("2", "PERCENT")
    # R10: two vs one year warranty in the same section is flagged
    war = [f for f in reqs if f["category"] == "Warranty"]
    assert len(war) == 2 and all(f["status"] == "CONFLICT" for f in war)
    # vector drawing: exact text, tokens with tags
    tokens = by_kind("token")
    assert {"AHU-1", "AHU-2"} <= {f["tag"] for f in tokens}
    assert any(f["category"] == "Steel shape" and f["value"] == "W12X26" for f in tokens)
    assert any(f["method"] == "PDF vector drawing text" for f in by_kind("drawing_text"))
    # schedule rows with quantities
    sched = by_kind("schedule_row")
    assert any(f["tag"] == "EF-1" and f["value"] == "2" for f in sched)
    # raster page went through drawing OCR
    assert any(f["page"] == 5 and f["method"] == "OCR (drawing, lines removed)" for f in by_kind("drawing_text"))

    wb = load_workbook(out)
    calc = wb["Calculations"]
    rows = list(calc.iter_rows(min_row=2, values_only=True))
    row = next(r for r in rows if "spare filter" in r[3])
    assert row[4] == 1 and row[7] == 2 and row[9].startswith("=E")       # formula, verified count only
    assert "AHU" in row[10]                                               # drawing evidence shown
    pages = {r[1]: r[2] for r in wb["Page Audit"].iter_rows(min_row=2, values_only=True)}
    assert pages == {1: "Spec text", 2: "Spec text", 3: "Drawing (vector PDF)", 4: "Spec text",
                     5: "Drawing (raster)"}
    assert Path(out.parent / "project_review" / "project_p0005.png").exists()
    report = json.loads(out.with_suffix(".run_report.json").read_text())
    assert report["conflicts_flagged"] == 2 and report["validated_cells"] > 100


def test_unverified_counts_are_never_multiplied(tmp_path):
    from spectakeoff.analysis import calculation_rows

    f = Finding(kind=REQUIREMENT, source_file="s.pdf", page=1, text="Provide two filters per pump.",
                method="x", value="two", unit="EACH", id="R1")
    (row,) = calculation_rows([f], {})
    assert row[7] is None and row[9] is False and "Needs verified count" in row[11]


def test_excel_keeps_formula_like_text_and_detects_tampering(tmp_path):
    f = Finding(kind=REQUIREMENT, source_file="s.pdf", page=1, text="=HYPERLINK(\"x\") shall\x07 be provided",
                method="PDF text layer", id="R1")
    res = SourceResult("s.pdf", "abc", [PageRecord("s.pdf", 1, "spec_text", "PDF text layer")], [f])
    out = tmp_path / "w.xlsx"
    build_workbook(out, [res], [], [], {}, "lv")
    wb = load_workbook(out)
    cell = wb["Prasības"]["G2"]
    assert cell.value == "=HYPERLINK(\"x\") shall be provided" and cell.data_type == "s"
    with pytest.raises(RuntimeError, match="cell mismatch"):
        _validate(out, {"Prasības": [[c.value for c in wb["Prasības"][1]], ["tampered"]]})


def test_evaluate_precision_recall():
    found = [{"source_file": "a.pdf", "page": 1, "kind": "drawing_text", "text": "W12x26 BEAM"},
             {"source_file": "a.pdf", "page": 1, "kind": "drawing_text", "text": "noise"},
             {"source_file": "a.pdf", "page": 2, "kind": "drawing_text", "text": "out of scope"}]
    gold = [{"source": "a.pdf", "page": "1", "kind": "drawing_text", "text": "W12x26 BEAM"},
            {"source": "a.pdf", "page": "1", "kind": "drawing_text", "text": "AHU-1"}]
    r = evaluate(found, gold)
    assert (r["recall"], r["precision"]) == (0.5, 0.5)


@pytest.mark.skipif("eng" not in installed_languages(), reason="Tesseract eng not installed")
def test_eval_cli_on_synthetic_image(tmp_path):
    img = tmp_path / "sheet.png"
    _, truth = make_drawing(img)
    out = tmp_path / "sheet.xlsx"
    assert main([str(img), "-o", str(out)]) == 0
    gold = tmp_path / "gold.csv"
    with open(gold, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["source", "page", "kind", "text"])
        for t in truth:
            w.writerow(["sheet.png", 1, "drawing_text", t["text"]])
    rep = tmp_path / "eval.json"
    assert main(["eval", str(out.with_suffix(".findings.json")), str(gold), "--report", str(rep)]) == 0
    res = json.loads(rep.read_text())
    assert res["recall"] >= 0.9 and res["precision"] >= 0.9, res["missed"]


SCANNED = ["SECTION 26 05 00 - COMMON WORK RESULTS FOR ELECTRICAL", "", "1.5 CLOSEOUT SUBMITTALS",
           "A. Submit two copies of operation and maintenance manuals for each panelboard.",
           "B. Provide spare fuses equal to 10 percent of each size installed.",
           "C. The warranty period shall be one year from substantial completion."]


@pytest.mark.skipif("eng" not in installed_languages(), reason="Tesseract eng not installed")
@pytest.mark.parametrize("rotate", [0, 90])
def test_scanned_spec_page_ocr_with_rotation(tmp_path, rotate):
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 28)
    im = Image.new("L", (1700, 2200), 255)
    d = ImageDraw.Draw(im)
    for i, line in enumerate(SCANNED):
        d.text((120, 150 + 60 * i), line, fill=20, font=font)
    src = tmp_path / "scan.png"
    im.rotate(rotate, expand=True).save(src)
    out = tmp_path / "scan.xlsx"
    assert main([str(src), "-o", str(out), "--no-overlays"]) == 0
    found = json.loads(out.with_suffix(".findings.json").read_text())["findings"]
    assert [f["text"] for f in found] == SCANNED[3:]
    assert all(f["section"] == "26 05 00" and f["division"] == "26" for f in found)
    assert [(f["value"], f["unit"]) for f in found] == [("two", "COPIES"), ("10", "PERCENT"), ("one", "YEAR")]

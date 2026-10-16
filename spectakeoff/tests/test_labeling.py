"""Gold-set labelling workbook -> human labels -> recall / precision."""
import json

import pymupdf
import pytest
from openpyxl import load_workbook

from spectakeoff.cli import main
from spectakeoff.evaluate import evaluate, match_score, read_gold, wilson
from spectakeoff.labeling import page_candidates, parse_pages

STAR = "*" * 74
PAGE1 = f"""SECTION 23 05 00 - COMMON WORK RESULTS FOR HVAC

A. Provide two copies of operation and maintenance manuals for each air handling unit.

B. The Contractor shall maintain the site free of debris.

{STAR}
NOTE: The Designer shall edit the warranty period below.
{STAR}

C. The warranty period for fans shall be two years from substantial completion.

D. Submit commissioning reports with the final payment request.
"""
PAGE2 = """SD-10 Operation and Maintenance Data

    Plumbing System; G, [_____]

SD-11 Closeout Submittals

    Record Drawings; G, [_____]

1.4 QUALITY ASSURANCE

Submit welder qualifications before welding starts.
"""
# what a careful reviewer would mark on these pages
TRUE = {"A. Provide two copies", "C. The warranty period", "D. Submit commissioning", "Plumbing System",
        "Record Drawings"}


def make_pdf(path):
    doc = pymupdf.open()
    for text in (PAGE1, PAGE2):
        doc.new_page().insert_text((40, 50), text, fontsize=9)
    doc.save(path)


def fill(wb_path, unclear=(), leave_blank=()):
    wb = load_workbook(wb_path)
    ws = wb["Etalons"]
    for row in ws.iter_rows(min_row=2):
        text = row[4].value
        if any(text.startswith(t) for t in leave_blank):
            continue
        row[5].value = "?" if any(text.startswith(t) for t in unclear) else (
            "Jā" if any(text.startswith(t) for t in TRUE) else "Nē")
    wb.save(wb_path)


def test_candidates_and_page_spec():
    cands = page_candidates(PAGE1)
    assert "A. Provide two copies of operation and maintenance manuals for each air handling unit." in cands
    assert not any(c.startswith("*") for c in cands)                     # asterisk rules are not rows
    assert "Plumbing System; G, [_____]" in page_candidates(PAGE2)
    assert parse_pages("3, 7-9,12") == [3, 7, 8, 9, 12]


def test_match_prefers_the_exact_row():
    assert match_score("Plumbing System; G, [_____]", "Plumbing System; G, [_____]") == 1.0
    assert match_score("Bolts", "Submit certificates for anchor bolts and nuts used on the project.") < 0.85
    gold = [{"source": "s.pdf", "page": 1, "text": "Plumbing System", "label": "no"},
            {"source": "s.pdf", "page": 1, "text": "Plumbing System; G, [_____]", "label": "yes"}]
    found = [{"source_file": "s.pdf", "page": 1, "kind": "requirement", "text": "Plumbing System; G, [_____]"}]
    r = evaluate(found, gold)
    assert (r["recall"], r["precision"]) == (1.0, 1.0)


def test_wilson_interval():
    lo, hi = wilson(9, 10)
    assert 0.55 < lo < 0.6 and 0.98 < hi < 0.99
    assert wilson(0, 0) is None


def test_label_then_eval_roundtrip(tmp_path):
    pdf = tmp_path / "spec.pdf"
    make_pdf(pdf)
    wb_path = tmp_path / "etalons.xlsx"
    assert main(["label", str(pdf), "--pages", "1-2", "-o", str(wb_path)]) == 0
    wb = load_workbook(wb_path)
    assert wb.sheetnames == ["Instrukcija", "Etalons", "Lapas", "_meta"]
    assert wb["_meta"].sheet_state == "hidden"
    assert wb["Etalons"].data_validations.dataValidation[0].formula1 == '"Jā,Nē,?"'
    assert wb["Etalons"].max_column == 7                           # no prediction column by default
    meta = json.loads(wb["_meta"]["A1"].value)
    assert meta["sources"][0]["name"] == "spec.pdf" and len(meta["sources"][0]["sha256"]) == 64

    fill(wb_path)
    rep = tmp_path / "eval.json"
    xlsx = tmp_path / "eval.xlsx"
    assert main(["eval", str(wb_path), "--report", str(rep), "--xlsx", str(xlsx)]) == 0
    res = json.loads(rep.read_text())
    assert res["mode"] == "full" and res["pages_scored"] == 2
    assert res["gold_requirements"] == 5
    assert (res["recall"], res["precision"]) == (1.0, 1.0), (res["missed"], res["false_positives"])
    assert load_workbook(xlsx).sheetnames == ["Kopsavilkums", "Palaistas garām", "Kļūdaini atrastas",
                                              "Pareizi atrastas"]


def test_misses_false_positives_unclear_and_incomplete_pages(tmp_path):
    pdf = tmp_path / "spec.pdf"
    make_pdf(pdf)
    wb_path = tmp_path / "etalons.xlsx"
    assert main(["label", str(pdf), "--pages", "1-2", "-o", str(wb_path)]) == 0
    # reviewer says welder qualifications ARE a requirement (tool does not extract it) -> a miss;
    # "Record Drawings" left blank -> page 2 is incomplete and not scored
    TRUE.add("Submit welder")
    try:
        fill(wb_path, unclear=("D. Submit commissioning",), leave_blank=("Record Drawings",))
    finally:
        TRUE.discard("Submit welder")
    gold, _ = read_gold(wb_path)
    found = json.loads("[]")
    res = evaluate(found, gold)
    assert res["pages_incomplete"] == ["spec.pdf p2"] and res["pages_scored"] == 1
    assert res["gold_unclear"] == 1 and res["gold_requirements"] == 2 and res["recall"] == 0.0

    wb = load_workbook(wb_path)
    wb["Etalons"]["F2"].value = "varbūt"
    wb.save(wb_path)
    with pytest.raises(ValueError, match="Nesaprotama"):
        read_gold(wb_path)
    assert main(["eval", str(wb_path)]) == 2


def test_eval_warns_and_fails_cleanly_when_source_missing(tmp_path):
    pdf = tmp_path / "spec.pdf"
    make_pdf(pdf)
    wb_path = tmp_path / "etalons.xlsx"
    assert main(["label", str(pdf), "--pages", "1", "-o", str(wb_path)]) == 0
    fill(wb_path)
    pdf.unlink()
    assert main(["eval", str(wb_path)]) == 2

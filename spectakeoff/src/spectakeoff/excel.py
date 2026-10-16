"""Workbook writer with full read-back validation (R12).

* strings are always stored as text, so a quote starting with "=" is never
  turned into a formula (v3 crashed / was formula-injectable here),
* XML-illegal control characters are removed (v3 raised IllegalCharacterError),
* after saving, every data sheet is reopened and compared cell by cell before
  the temporary file is renamed to the final name.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from spectakeoff import __version__
from spectakeoff.i18n import Localizer
from spectakeoff.models import (CONFLICT, DRAWING_TEXT, REQUIREMENT, REVIEW, SCHEDULE_ROW, TOKEN,
                                VISUAL_REVIEW, Finding, SourceResult)

FINDING_COLS = ["col.id", "col.source", "col.page", "col.region", "col.kind", "col.category", "col.text",
                "col.value", "col.unit", "col.tag", "col.section", "col.section_title", "col.division",
                "col.method", "col.confidence", "col.status", "col.flags", "col.translated", "col.sha"]
WIDTHS = [24, 26, 7, 20, 14, 26, 70, 14, 10, 12, 14, 30, 9, 26, 11, 14, 60, 40, 24]

HEADER_FILL = PatternFill("solid", fgColor="17365D")
FILLS = {"REVIEW": PatternFill("solid", fgColor="FCE4D6"), "CONFLICT": PatternFill("solid", fgColor="F4CCCC"),
         "CANDIDATE": PatternFill("solid", fgColor="FFF2CC")}

RULES = [
    ("R1", "Every row is supported by PDF text, OCR or table evidence.",
     "IMPLEMENTED — verbatim evidence, page and region stored for every finding."),
    ("R2", "Evidence must match the cited page.", "PARTIAL — evidence comes from that page; no semantic re-check."),
    ("R3", "Spec section must exist and match the page.",
     "PARTIAL — stricter heading rules than v3; still heuristic, sections not found are flagged."),
    ("R4", "Quantity and unit come from the source text.", "IMPLEMENTED — textual quantity only, incl. percent."),
    ("R5", "Unspecified quantity is never invented.", "IMPLEMENTED — blank + flag."),
    ("R6", "No duplicate rows or padding.", "PARTIAL — exact duplicates per page and IoU duplicates per drawing removed."),
    ("R7", "Never force a target row count.", "IMPLEMENTED."),
    ("R8", "Per-unit quantities multiply only verified counts.",
     "IMPLEMENTED — Calculations sheet multiplies only estimator-verified counts (--counts CSV); drawing/schedule "
     "counts are shown as unverified evidence."),
    ("R9", "Labor uses explicit factors only.", "NOT IMPLEMENTED in 4.0.0."),
    ("R10", "Conflicting requirements are flagged.",
     "PARTIAL — different quantities for the same section+category+unit are flagged CONFLICT."),
    ("R11", "Pages with signals but no rows, and unreadable pages, need review.",
     "IMPLEMENTED — page audit, section coverage, unreadable drawing labels and photos go to review."),
    ("R12", "Reopen Excel and validate saved output.", "IMPLEMENTED — every cell of every data sheet compared."),
    ("R13", "OCR scanned pages; low confidence flagged.",
     "IMPLEMENTED — page OCR with rotation detection; drawing OCR with line removal and noise filter."),
    ("R14", "Drawing quantities need source tags/schedules or estimator counts.",
     "PARTIAL — tags, count notes and schedules extracted with evidence; no automatic counting from images."),
    ("R15", "Source quote is never translated or reworded.", "IMPLEMENTED — translation is a separate column."),
    ("R16", "Division 23/26 coverage checks.",
     "PARTIAL — Division 23/26 sections without candidates are flagged in Section Coverage."),
]

METHOD = [
    ("Spec text", "PDF text layer first; pages with too little text are rendered and OCR'd with Tesseract "
                  "(orientation detected via OSD)."),
    ("Vector drawings", "Text is read directly from the PDF drawing (exact, no OCR) with its position; "
                        "construction tokens (steel shapes, bolts, tags, dimensions, references, count notes) "
                        "are recognised from that text."),
    ("Raster drawings", "Text height is estimated and the image upscaled; long horizontal, vertical and "
                        "diagonal lines are removed; text-line regions are OCR'd at 0° and 90°; every word "
                        "passes a noise filter (vocabulary, token patterns, confidence)."),
    ("Unreadable labels", "Regions that look like a text line but cannot be read reliably are listed as "
                          "visual-review findings with their position; they are never guessed."),
    ("Photographs", "Photos are detected and sent to visual review. 4.0 does not count objects (bolts, members) "
                    "from photos: classical image processing was measured as unreliable for this."),
    ("Schedules", "Tables are detected on PDF pages with ruling lines; each row is kept with its header."),
    ("Calculations", "Totals are Excel formulas over per-unit quantity × estimator-verified count only."),
    ("Confidence", "OCR confidence is Tesseract word confidence — not a measure of correctness."),
    ("Accuracy", "Nothing is certified. Measure precision/recall with `spectakeoff eval` against a "
                 "human-checked gold set before any accuracy claim."),
    ("Privacy", "Optional --translate sends text to an external service. Off by default; do not enable for "
                "confidential documents without authorisation."),
]


def clean(v: Any) -> Any:
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def finding_row(f: Finding, t: Localizer) -> list[Any]:
    return [f.id, f.source_file, f.page, str(f.bbox) if f.bbox else "", t("kind." + f.kind), f.category, f.text,
            f.value, f.unit, f.tag, f.section, f.section_title, f.division, f.method,
            None if f.confidence is None else round(f.confidence, 1), t("status." + f.status),
            "; ".join(f.flags), f.translated, f.page_sha256]


class _Sheets:
    def __init__(self, wb: Workbook):
        self.wb = wb
        self.expected: dict[str, list[list[Any]]] = {}

    def add(self, title: str, header: Sequence[str], rows: Sequence[Sequence[Any]],
            widths: Sequence[int] | None = None, status_col: int | None = None):
        ws = self.wb.create_sheet(title[:31])
        data = [[clean(v) for v in header]] + [[clean(v) for v in r] for r in rows]
        for r in data:
            ws.append([None if v == "" else v for v in r])
        for row in ws.iter_rows(min_row=1):
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = "s"     # never interpret evidence text as a formula
        _style(ws, widths)
        if status_col is not None:
            for i in range(2, ws.max_row + 1):
                key = next((k for k in FILLS if str(ws.cell(i, status_col).value or "") in _STATUS_ALIASES[k]), None)
                if key:
                    ws.cell(i, status_col).fill = FILLS[key]
        self.expected[ws.title] = [[None if v == "" else v for v in r] for r in data]
        return ws


_STATUS_ALIASES = {"CANDIDATE": ("CANDIDATE", "KANDIDĀTS"), "REVIEW": ("REVIEW", "PĀRBAUDĪT"),
                   "CONFLICT": ("CONFLICT", "PRETRUNA")}


def _style(ws, widths=None):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(1, ws.max_column + 1):
        w = widths[i - 1] if widths and i <= len(widths) else 18
        ws.column_dimensions[get_column_letter(i)].width = w


def _priority(f: Finding) -> int:
    if f.status == CONFLICT:
        return 1
    if f.kind == VISUAL_REVIEW:
        return 2
    if f.status == REVIEW:
        return 3
    return 4


def build_workbook(out: Path, results: list[SourceResult], calc_rows: list[list[Any]],
                   coverage_rows: list[list[Any]], run_meta: dict, lang: str = "en") -> dict:
    t = Localizer(lang)
    wb = Workbook()
    wb.remove(wb.active)
    S = _Sheets(wb)
    findings = [f for r in results for f in r.findings]
    pages = [p for r in results for p in r.pages]
    hdr = [t(c) for c in FINDING_COLS]

    def by(kind):
        return [finding_row(f, t) for f in findings if f.kind == kind]

    status_col = FINDING_COLS.index("col.status") + 1
    summary = [
        ("Tool version", __version__),
        ("Sources", ", ".join(r.source for r in results)),
        ("Sources failed", ", ".join(f"{r.source}: {r.error}" for r in results if r.error) or "0"),
        ("Pages / images", len(pages)),
        ("Pages requiring review", sum(1 for p in pages if p.status == "REVIEW")),
        ("Spec requirement candidates", sum(1 for f in findings if f.kind == REQUIREMENT)),
        ("Drawing text lines", sum(1 for f in findings if f.kind == DRAWING_TEXT)),
        ("Drawing tokens", sum(1 for f in findings if f.kind == TOKEN)),
        ("Schedule rows", sum(1 for f in findings if f.kind == SCHEDULE_ROW)),
        ("Visual review items", sum(1 for f in findings if f.kind == VISUAL_REVIEW)),
        ("Conflicts flagged", sum(1 for f in findings if f.status == CONFLICT)),
        ("Rows certified correct", 0),
        ("Accuracy claim", "Not certified; every row requires human review."),
        ("Delivery gate", "Do not deliver as a final takeoff until the review queue and source pages are checked."),
    ] + [(r.source + " SHA-256", r.sha256) for r in results]
    S.add(t("sheet.summary"), [t("col.metric"), t("col.v")], summary, [34, 110])
    S.add(t("sheet.requirements"), hdr, by(REQUIREMENT), WIDTHS, status_col)
    S.add(t("sheet.drawing_text"), hdr, by(DRAWING_TEXT), WIDTHS, status_col)
    S.add(t("sheet.tokens"), hdr, by(TOKEN), WIDTHS, status_col)
    S.add(t("sheet.schedules"), hdr, by(SCHEDULE_ROW), WIDTHS, status_col)
    calc_hdr = [t("col.req_id"), t("col.source"), t("col.page"), t("col.text"), t("col.per_unit"), t("col.unit"),
                t("col.counted_item"), t("col.verified_count"), t("col.count_source"), t("col.total"),
                t("col.drawing_evidence"), t("col.status")]
    ws_calc = S.add(t("sheet.calc"), calc_hdr, [r[:9] + [""] + r[10:] for r in calc_rows],
                    [24, 24, 7, 60, 12, 10, 18, 12, 30, 14, 50, 30])
    for i, r in enumerate(calc_rows, start=2):   # formula only where a verified count exists
        if r[9]:
            ws_calc.cell(i, 10).value = f"=E{i}*H{i}"
            S.expected[ws_calc.title][i - 1][9] = f"=E{i}*H{i}"
    review = sorted(findings, key=lambda f: (_priority(f), f.source_file, f.page))
    S.add(t("sheet.review"), [t("col.priority")] + hdr, [[_priority(f)] + finding_row(f, t) for f in review],
          [8] + WIDTHS, status_col + 1)
    page_hdr = [t("col.source"), t("col.page"), t("col.page_kind"), t("col.method"), t("col.chars"),
                t("col.images"), t("col.confidence"), t("col.section"), t("col.section_title"), t("col.regions"),
                t("col.accepted"), t("col.rejected"), t("col.findings"), t("col.status"), t("col.warnings"),
                t("col.overlay"), t("col.sha")]
    page_rows = [[p.source_file, p.page, t("page." + p.page_kind), p.method, p.characters, p.images,
                  None if p.ocr_confidence is None else round(p.ocr_confidence, 1), p.section, p.section_title,
                  p.text_regions, p.accepted_lines, p.rejected_noise, p.findings,
                  t("status.REVIEW") if p.status == "REVIEW" else p.status, "; ".join(p.warnings), p.overlay,
                  p.sha256] for p in pages]
    S.add(t("sheet.pages"), page_hdr, page_rows, [24, 7, 20, 26, 10, 9, 10, 14, 30, 9, 9, 9, 9, 12, 60, 40, 24])
    S.add(t("sheet.coverage"), [t("col.section"), t("col.section_title"), t("col.division"), t("col.pages"),
                                t("col.hits"), t("col.findings"), t("col.coverage")], coverage_rows,
          [16, 40, 9, 14, 14, 10, 60])
    S.add(t("sheet.rules"), [t("col.rule"), t("col.required"), t("col.impl")], RULES, [8, 60, 90])
    S.add(t("sheet.method"), [t("col.topic"), t("col.details")],
          METHOD + [("Run metadata", json.dumps(run_meta, ensure_ascii=False, default=str))], [22, 130])

    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(out.stem + ".tmp.xlsx")
    wb.save(tmp)
    try:
        _validate(tmp, S.expected)
    except Exception:
        tmp.unlink(missing_ok=True)
        raise
    tmp.replace(out)
    return {"output": str(out), "sheets": list(S.expected), "validated_cells": sum(
        len(r) for rows in S.expected.values() for r in rows)}


def _validate(path: Path, expected: dict[str, list[list[Any]]]) -> None:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        for name, rows in expected.items():
            if name not in wb.sheetnames:
                raise RuntimeError(f"Excel validation: missing sheet {name!r}")
            got = [list(r) for r in wb[name].iter_rows(values_only=True)]
            if len(got) != len(rows):
                raise RuntimeError(f"Excel validation: {name!r} has {len(got)} rows, expected {len(rows)}")
            for i, (g, e) in enumerate(zip(got, rows), start=1):
                g = (g + [None] * len(e))[:len(e)]
                if any(not _same(a, b) for a, b in zip(g, e)):
                    raise RuntimeError(f"Excel validation: cell mismatch in {name!r} row {i}")
    finally:
        wb.close()


def _same(a: Any, b: Any) -> bool:
    if isinstance(a, float) or isinstance(b, float):
        try:
            return abs(float(a) - float(b)) < 1e-9
        except (TypeError, ValueError):
            return False
    return a == b

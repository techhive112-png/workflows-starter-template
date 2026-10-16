"""Gold-set labelling workbook for measuring precision AND recall.

`spectakeoff label` lists every sentence / list item on a sample of pages. A person
marks each one Jā/Nē ("is this a closeout requirement?"). Because every candidate
on those pages is labelled - not only what the tool found - `spectakeoff eval`
can measure recall (what the tool missed) as well as precision.
"""
from __future__ import annotations

import dataclasses
import hashlib
import json
import random
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from spectakeoff import __version__
from spectakeoff.excel import clean
from spectakeoff.specs import split_spans

SHEET = "Etalons"
META = "_meta"
HEAD = ["ID", "Fails", "Lapa", "Nr. lapā", "Teksts (oriģināls)", "Prasība? (Jā/Nē)", "Piezīme"]
PRED_HEAD = "Programma atrada"

INSTRUCTIONS = [
    ("Mērķis", "Atzīmē KATRU rindu šajās lapās: vai tā ir nodošanas / closeout prasība? Tad "
               "`spectakeoff eval` izmērīs, ko programma atrod pareizi, ko kļūdaini un ko palaiž garām."),
    ("Kā aizpildīt", "Kolonnā F izvēlies Jā vai Nē (vai ? ja tiešām nav skaidrs - tādas rindas netiek "
                     "vērtētas). Atver PDF norādītajā lapā, ja teksts nav saprotams bez konteksta."),
    ("Jā", "Teikums vai saraksta punkts uzliek darbuzņēmējam pienākumu NODOT vai IESNIEGT kaut ko "
           "nodošanai/closeout: garantija, O&M rokasgrāmatas, izpilddokumentācija (as-built / record "
           "drawings), rezerves daļas, papildu materiāli (attic stock), speciālie instrumenti, apmācība, "
           "testu atskaites, sertifikāti un atbilstības apliecinājumi, iesniegumu saraksta punkti "
           "(piem. UFGS SD-06/07/09/10/11)."),
    ("Nē", "Būvdarbu metodes, materiālu īpašības, juridiski skaidrojumi ('shall not relieve...'), "
           "redaktora piezīmes (NOTE:, 'The Engineer shall...'), virsraksti, lapu galvenes, atsauces uz "
           "citām sadaļām, uzstādīšana 'saskaņā ar ražotāja instrukcijām'."),
    ("Piemērs Jā", "Provide two (2) copies of all manufacturers warranties specified for materials."),
    ("Piemērs Jā", "Operation and Maintenance Manuals; G, [_____]   (punkts zem SD-10)"),
    ("Piemērs Nē", "Air-entraining admixtures shall meet the requirements of ASTM C260."),
    ("Piemērs Nē", "Manufacturer's certifications shall not relieve the Contractor of the responsibility..."),
    ("Svarīgi", "Neskaties, ko atrada programma, pirms neesi atzīmējis pats - citādi mērījums būs "
                "neobjektīvs. Ja failā ir kolonna 'Programma atrada', to var paslēpt."),
    ("Pēc tam", "Saglabā failu un palaid:  spectakeoff eval etalons.xlsx   (programma pati atkārtoti "
                "apstrādās tos pašus PDF un salīdzinās). Lapas, kur kaut viena rinda nav atzīmēta, "
                "netiek vērtētas."),
]


def _n(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").replace("’", "'")).strip().casefold()


def page_candidates(text: str) -> list[str]:
    """Every sentence and list item on a page, in order, without duplicates."""
    spans: list[str] = []
    for span in split_spans(text):
        span = re.sub(r"^\*{5,}\s*|\s*\*{5,}$", "", span.strip()).strip()   # asterisk note rules
        if not span or set(span) <= set("*[] ."):
            continue
        if sum(c.isalnum() for c in span) < 2 and spans:    # "G, [_____]" belongs to the item before
            spans[-1] += " " + span
            continue
        spans.append(span)
    out, seen = [], set()
    for span in spans:
        key = _n(span)
        if key not in seen:
            seen.add(key)
            out.append(span)
    return out


def parse_pages(spec: str) -> list[int]:
    pages: set[int] = set()
    for part in (spec or "").split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            pages.update(range(int(a), int(b) + 1))
        else:
            pages.add(int(part))
    return sorted(pages)


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build_label_workbook(sources: list[Path], out: Path, cfg, sample: int = 20, pages: str = "",
                         seed: int = 1, show_predictions: bool = False) -> dict:
    from spectakeoff.evaluate import match_score
    from spectakeoff.pipeline import process_source

    cfg = dataclasses.replace(cfg, keep_text=True, overlays=False)
    results = [process_source(p, cfg)[0] for p in sources]
    eligible = [(r, rec) for r in results for rec in r.pages if rec.text.strip()]
    if pages:
        wanted = set(parse_pages(pages))
        chosen = [(r, rec) for r, rec in eligible if rec.page in wanted]
    else:
        rng = random.Random(seed)
        chosen = sorted(rng.sample(eligible, min(sample, len(eligible))),
                        key=lambda x: (results.index(x[0]), x[1].page))
    rows = []
    for r, rec in chosen:
        preds = [f for f in r.findings if f.kind == "requirement" and f.page == rec.page]
        for k, cand in enumerate(page_candidates(rec.text), 1):
            found = any(match_score(f.text, cand) >= 0.85 for f in preds)
            rows.append([f"{Path(r.source).stem[:12]}-P{rec.page:04d}-{k:03d}", r.source, rec.page, k, cand,
                         None, None] + (["Jā" if found else "Nē"] if show_predictions else []))

    wb = Workbook()
    ins = wb.active
    ins.title = "Instrukcija"
    ins.append(["Tēma", "Apraksts"])
    for row in INSTRUCTIONS:
        ins.append(list(row))
    ins.column_dimensions["A"].width = 16
    ins.column_dimensions["B"].width = 120
    for row in ins.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ins[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="17365D")

    ws = wb.create_sheet(SHEET)
    ws.append(HEAD + ([PRED_HEAD] if show_predictions else []))
    for row in rows:
        ws.append([clean(v) for v in row])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str):
                c.data_type = "s"
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="17365D")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for col, w in zip("ABCDEFGH", (22, 22, 7, 8, 90, 16, 30, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
    last = max(2, ws.max_row)
    dv = DataValidation(type="list", formula1='"Jā,Nē,?"', allow_blank=True,
                        error="Izvēlies Jā, Nē vai ?", errorTitle="SpecTakeoff")
    ws.add_data_validation(dv)
    dv.add(f"F2:F{last}")
    ws.conditional_formatting.add(f"F2:F{last}", FormulaRule(formula=['F2="Jā"'],
                                                            fill=PatternFill("solid", fgColor="D9EAD3")))
    ws.conditional_formatting.add(f"F2:F{last}", FormulaRule(formula=['F2="Nē"'],
                                                            fill=PatternFill("solid", fgColor="F4CCCC")))
    ws.conditional_formatting.add(f"F2:F{last}", FormulaRule(formula=['F2=""'],
                                                            fill=PatternFill("solid", fgColor="FFF2CC")))

    pg = wb.create_sheet("Lapas")
    pg.append(["Fails", "Lapa", "Rindas", "Atzīmētas", "Gatava?"])
    for i, (r, rec) in enumerate(chosen, start=2):
        pg.append([r.source, rec.page, None, None, None])
        pg.cell(i, 3).value = f'=COUNTIFS({SHEET}!B:B,A{i},{SHEET}!C:C,B{i})'
        pg.cell(i, 4).value = f'=COUNTIFS({SHEET}!B:B,A{i},{SHEET}!C:C,B{i},{SHEET}!F:F,"<>")'
        pg.cell(i, 5).value = f'=IF(D{i}=C{i},"Jā","Nē")'
    for col, w in zip("ABCDE", (28, 8, 10, 12, 10)):
        pg.column_dimensions[col].width = w
    for c in pg[1]:
        c.font = Font(bold=True)

    meta = wb.create_sheet(META)
    meta.sheet_state = "hidden"
    info = {"tool_version": __version__, "seed": seed, "sample": sample, "pages_spec": pages,
            "show_predictions": show_predictions,
            "sources": [{"name": p.name, "path": str(p.resolve()), "sha256": _sha(p)} for p in sources],
            "pages": [{"source": r.source, "page": rec.page} for r, rec in chosen]}
    meta["A1"] = json.dumps(info, ensure_ascii=False)
    meta["A1"].data_type = "s"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"output": str(out), "pages": len(chosen), "rows": len(rows),
            "eligible_pages": len(eligible)}

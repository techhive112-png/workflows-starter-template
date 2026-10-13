#!/usr/bin/env python3
"""
SpecTakeoff batch - run several specifications and compare them on one sheet.

This is the tool for open task 5.2 (testing more public specs). It processes each PDF, runs the
self-test on each, and writes one Batch_Summary.xlsx plus a table on screen. It never merges rows
from different specs: each spec keeps its own schedule file.

  python batch.py a.pdf b.pdf c.pdf        run these
  python batch.py --fetch urls.txt         download each URL (one per line) first, then run
  python batch.py *.pdf --no-selftest      faster, skips the per-spec quality gate

A spec that fails the self-test is still reported, marked FAILED, with the reason. Nothing is
hidden: the point of a batch is to see which spec formats break the extractor.
"""
import argparse
import os
import re
import subprocess
import sys
import time
import urllib.parse
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
COLS = ["Spec", "Pages", "Rows", "Numbering", "Coverage sheet", "Sections with rows",
        "Agent errors", "Agent warnings", "Seconds", "Self-test", "Types found", "Notes"]


def fetch(url, into):
    name = os.path.basename(urllib.parse.urlparse(url).path) or "spec.pdf"
    name = urllib.parse.unquote(name)
    if not name.lower().endswith(".pdf"):
        name += ".pdf"
    dest = os.path.join(into, re.sub(r"[^A-Za-z0-9._-]+", "_", name))
    if os.path.exists(dest):
        print(f"  already downloaded: {os.path.basename(dest)}")
        return dest
    print(f"  downloading {url}")
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(req, timeout=300) as r, open(dest, "wb") as f:
            f.write(r.read())
    except Exception as e:                                           # noqa: BLE001
        print(f"  COULD NOT DOWNLOAD: {e}")
        return None
    if open(dest, "rb").read(5) != b"%PDF-":
        print("  NOT A PDF (the site probably returned an error page)")
        os.remove(dest)
        return None
    return dest


def one(pdf, selftest=True):
    from openpyxl import load_workbook
    name = os.path.basename(pdf)
    out = os.path.join(os.path.dirname(os.path.abspath(pdf)),
                       re.sub(r"\.pdf$", "", name, flags=re.I) + "_Closeout_Schedule.xlsx")
    t0 = time.time()
    p = subprocess.run([sys.executable, os.path.join(HERE, "spectakeoff.py"), pdf, "-o", out],
                       capture_output=True, text=True)
    secs = time.time() - t0
    log = p.stdout + p.stderr
    row = {c: "" for c in COLS}
    row["Spec"], row["Seconds"] = name, round(secs)
    m = re.search(r"\[extract\]\s+(\d+) pages", log)
    row["Pages"] = int(m.group(1)) if m else ""
    if not os.path.exists(out):
        row["Rows"], row["Notes"] = 0, "NOT DELIVERED - " + (log.strip().splitlines() or ["no output"])[-1][:200]
        row["Self-test"] = "n/a"
        return row, out
    wb = load_workbook(out)
    ws = wb["Closeout Schedule"]
    rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if any(v not in (None, "") for v in r)]
    row["Rows"] = len(rows)
    secs_col = [str(r[0]) for r in rows]
    row["Numbering"] = ("CSI 1995 five-digit" if any(" " not in s and s.isdigit() for s in secs_col)
                        else "MasterFormat 2004")
    if any(" " not in s and s.isdigit() for s in secs_col) and any(" " in s for s in secs_col):
        row["Numbering"] = "both (mixed manual)"
    cov = [s for s in wb.sheetnames if "Coverage" in s]
    row["Coverage sheet"] = cov[0] if cov else "-"
    row["Sections with rows"] = len({str(r[0]) for r in rows})
    ch = [r for r in wb["Checks"].iter_rows(min_row=2, values_only=True)]
    row["Agent errors"] = sum(1 for r in ch if r[1] == "ERROR")
    row["Agent warnings"] = sum(1 for r in ch if r[1] == "WARN")
    from collections import Counter
    row["Types found"] = ", ".join(f"{t} {n}" for t, n in Counter(str(r[2]) for r in rows).most_common())
    if selftest:
        q = subprocess.run([sys.executable, os.path.join(HERE, "selftest.py"), pdf],
                           capture_output=True, text=True)
        row["Self-test"] = "PASS" if q.returncode == 0 else "FAILED"
        if q.returncode != 0:
            bad = [l for l in q.stdout.splitlines() if l.startswith("[FAIL]")]
            row["Notes"] = "; ".join(b[6:].strip() for b in bad)[:300]
    else:
        row["Self-test"] = "skipped"
    return row, out


def main():
    ap = argparse.ArgumentParser(description="SpecTakeoff batch runner")
    ap.add_argument("pdfs", nargs="*")
    ap.add_argument("--fetch", help="file of URLs, one per line")
    ap.add_argument("--no-selftest", action="store_true")
    ap.add_argument("-o", "--out", default=os.path.join(HERE, "Batch_Summary.xlsx"))
    a = ap.parse_args()

    pdfs = list(a.pdfs)
    if a.fetch:
        into = os.path.join(os.path.dirname(os.path.abspath(a.fetch)), "specs")
        os.makedirs(into, exist_ok=True)
        print("fetching:")
        for url in [l.strip() for l in open(a.fetch, encoding="utf-8") if l.strip()
                    and not l.strip().startswith("#")]:
            got = fetch(url, into)
            if got:
                pdfs.append(got)
    if not pdfs:
        sys.exit("No PDF to process. Give PDF paths, or --fetch a file of URLs.")

    results = []
    for pdf in pdfs:
        print(f"\n=== {os.path.basename(pdf)} ===")
        r, out = one(pdf, not a.no_selftest)
        results.append(r)
        print(f"  {r['Rows']} rows, {r['Pages']} pages, {r['Seconds']}s, self-test {r['Self-test']}"
              + (f"  >> {r['Notes']}" if r["Notes"] else ""))

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Batch Summary"
    ws.append(COLS)
    for r in results:
        ws.append([r[c] for c in COLS])
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if str(row[9].value) != "PASS":
            for c in row:
                c.fill = PatternFill("solid", fgColor="FFF2CC")
    for i, w in enumerate([34, 7, 7, 20, 20, 10, 9, 10, 9, 10, 46, 60]):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(a.out)

    print("\n%-34s %5s %5s %-20s %-9s %s" % ("SPEC", "PAGES", "ROWS", "NUMBERING", "SELF-TEST", "SECONDS"))
    for r in results:
        print("%-34s %5s %5s %-20s %-9s %s" % (r["Spec"][:34], r["Pages"], r["Rows"],
                                               r["Numbering"], r["Self-test"], r["Seconds"]))
    print(f"\nSummary written to {a.out}")
    sys.exit(0 if all(r["Self-test"] in ("PASS", "skipped") for r in results) else 1)


if __name__ == "__main__":
    main()
